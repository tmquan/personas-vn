"""Embed figures into Nemotron Personas Vietnam v1.pptx.

Slide layout: 40 x 22.5 inches. Title bar 1.0-3.7, subtitle 3.7-4.7,
body 6.5-20.5. We narrow body text placeholders on slides that get
images, then add pictures in the right-side empty area. Text content
itself is never modified.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

DECK = Path("Nemotron Personas Vietnam v1.pptx")
FIG_MAPS = Path("docs/figures/maps")
FIG_ANALYSIS = Path("docs/figures/analysis")


def find_shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"shape {name!r} not found on slide")


def resize_shape_width(shape, new_width_in: float) -> None:
    """Narrow the shape's width in place; keep left/top/height."""
    shape.width = Inches(new_width_in)


def add_picture_fixed(slide, path: Path, left_in, top_in, width_in, height_in):
    """Insert picture, returning the inserted shape."""
    pic = slide.shapes.add_picture(
        str(path),
        left=Inches(left_in),
        top=Inches(top_in),
        width=Inches(width_in),
        height=Inches(height_in),
    )
    return pic


def main() -> None:
    p = Presentation(str(DECK))
    log: list[tuple[int, str, tuple[float, float, float, float]]] = []

    # ---- Slide 1: Title — country reference (right-half, watermark feel) ----
    s = p.slides[0]
    fig = FIG_MAPS / "01_country_reference.png"
    coords = (22.0, 5.5, 14.0, 14.0 / 1.222)
    add_picture_fixed(s, fig, *coords)
    log.append((1, str(fig), coords))

    # ---- Slide 4: Layer 1 — analysis/01 + maps/01 ----
    s = p.slides[3]
    body = find_shape_by_name(s, "Text Placeholder 1")
    resize_shape_width(body, 17.0)
    fig_a = FIG_ANALYSIS / "01_tables_per_database.png"
    coords_a = (19.0, 6.5, 19.0, 19.0 / 1.964)
    add_picture_fixed(s, fig_a, *coords_a)
    log.append((4, str(fig_a), coords_a))
    fig_b = FIG_MAPS / "01_country_reference.png"
    coords_b = (29.0, 16.5, 6.0, 6.0 / 1.222)
    add_picture_fixed(s, fig_b, *coords_b)
    log.append((4, str(fig_b), coords_b))

    # ---- Slide 5: Layer 2 — coverage matrix ----
    s = p.slides[4]
    body = find_shape_by_name(s, "Text Placeholder 1")
    resize_shape_width(body, 17.0)
    fig = FIG_ANALYSIS / "38_coverage_matrix.png"
    width = 19.0
    height = width / 2.097
    top = 6.5 + (14.0 - height) / 2
    coords = (19.0, top, width, height)
    add_picture_fixed(s, fig, *coords)
    log.append((5, str(fig), coords))

    # ---- Slide 6: Layer 3 — persona provenance Sankey ----
    s = p.slides[5]
    body = find_shape_by_name(s, "Text Placeholder 1")
    resize_shape_width(body, 17.0)
    fig = FIG_ANALYSIS / "39_persona_provenance.png"
    width = 19.0
    height = width / 1.935
    top = 6.5 + (14.0 - height) / 2
    coords = (19.0, top, width, height)
    add_picture_fixed(s, fig, *coords)
    log.append((6, str(fig), coords))

    # ---- Slide 10: 4 parquets — population map (right, smaller) ----
    s = p.slides[9]
    body = find_shape_by_name(s, "Text Placeholder 1")
    resize_shape_width(body, 22.0)
    fig = FIG_MAPS / "02_population_by_province.png"
    width = 13.0
    height = width / 1.222
    top = 6.5 + (14.0 - height) / 2
    coords = (25.0, top, width, height)
    add_picture_fixed(s, fig, *coords)
    log.append((10, str(fig), coords))

    # ---- Slide 13: Where we are today — 2 stats stacked ----
    s = p.slides[12]
    body = find_shape_by_name(s, "Text Placeholder 1")
    resize_shape_width(body, 17.0)
    fig_top = FIG_ANALYSIS / "16_unemployment_by_education.png"
    width = 17.0
    height_top = width / 2.115
    coords_top = (20.0, 6.5, width, height_top)
    add_picture_fixed(s, fig_top, *coords_top)
    log.append((13, str(fig_top), coords_top))
    fig_bot = FIG_ANALYSIS / "13_occupation_structure_shift.png"
    height_bot = width / 2.258
    coords_bot = (20.0, 6.5 + height_top + 0.2, width, height_bot)
    add_picture_fixed(s, fig_bot, *coords_bot)
    log.append((13, str(fig_bot), coords_bot))

    # ---- Slide 14: Sample persona — country reference (small, right) ----
    s = p.slides[13]
    body = find_shape_by_name(s, "Text Placeholder 1")
    resize_shape_width(body, 25.0)
    fig = FIG_MAPS / "01_country_reference.png"
    width = 10.0
    height = width / 1.222
    top = 6.5 + (14.0 - height) / 2
    coords = (28.0, top, width, height)
    add_picture_fixed(s, fig, *coords)
    log.append((14, str(fig), coords))

    p.save(str(DECK))

    print("=== Insertion log ===")
    for slide_idx, fig_path, (l, t, w, h) in log:
        print(
            f"  slide {slide_idx:>2}  +  {fig_path}"
            f"  @ left={l:.2f}in top={t:.2f}in width={w:.2f}in height={h:.2f}in"
        )


if __name__ == "__main__":
    main()
